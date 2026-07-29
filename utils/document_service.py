"""
utils/document_service.py

Handles uploaded customer documents: extracts text (native for
.txt/.csv, best-effort optional-dependency support for .pdf/.docx/
.xlsx), then produces a summary and key-point extraction using a
dependency-free extractive method (frequency-scored sentence
ranking). This always works, even with zero NLP libraries installed;
swapping in a transformer summarizer later only touches
``summarize()``.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass

logger = logging.getLogger("ics.documents")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx", ".png", ".jpg", ".jpeg"}
MAX_UPLOAD_BYTES = 15 * 1024 * 1024  # 15 MB


@dataclass
class ExtractionResult:
    success: bool
    text: str = ""
    error: str = ""


def extract_text(filename: str, file_bytes: bytes) -> ExtractionResult:
    """Extract raw text from an uploaded file based on its extension."""

    if len(file_bytes) > MAX_UPLOAD_BYTES:
        return ExtractionResult(False, error="File exceeds the 15 MB upload limit.")

    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        return ExtractionResult(False, error=f"Unsupported file type: {ext or 'unknown'}")

    try:
        if ext == ".txt":
            return ExtractionResult(True, text=file_bytes.decode("utf-8", errors="ignore"))

        if ext == ".csv":
            text_io = io.StringIO(file_bytes.decode("utf-8", errors="ignore"))
            rows = list(csv.reader(text_io))
            preview = "\n".join(", ".join(row) for row in rows[:200])
            return ExtractionResult(True, text=preview)

        if ext == ".pdf":
            try:
                from pypdf import PdfReader  # type: ignore

                reader = PdfReader(io.BytesIO(file_bytes))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                return ExtractionResult(True, text=text)
            except ImportError:
                return ExtractionResult(False, error="PDF extraction requires the 'pypdf' package (pip install pypdf).")

        if ext == ".docx":
            try:
                import docx  # type: ignore

                document = docx.Document(io.BytesIO(file_bytes))
                text = "\n".join(p.text for p in document.paragraphs)
                return ExtractionResult(True, text=text)
            except ImportError:
                return ExtractionResult(False, error="DOCX extraction requires the 'python-docx' package.")

        if ext == ".xlsx":
            try:
                import openpyxl  # type: ignore

                workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
                lines = []
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True, max_row=200):
                        lines.append(", ".join(str(c) for c in row if c is not None))
                return ExtractionResult(True, text="\n".join(lines))
            except ImportError:
                return ExtractionResult(False, error="XLSX extraction requires the 'openpyxl' package.")

        if ext in {".png", ".jpg", ".jpeg"}:
            from utils.ocr_service import extract_text_from_image

            ocr_result = extract_text_from_image(file_bytes)
            if not ocr_result.success:
                return ExtractionResult(False, error=ocr_result.error)
            return ExtractionResult(True, text=ocr_result.text)

    except Exception as exc:  # pragma: no cover - defensive
        logger.exception("Text extraction failed for %s", filename)
        return ExtractionResult(False, error=f"Failed to read file: {exc}")

    return ExtractionResult(False, error="Unhandled file type.")


_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+")
_WORD_RE = re.compile(r"[a-zA-Z']+")

_COMMON_WORDS = {
    "the", "a", "an", "and", "or", "of", "to", "in", "on", "for", "is",
    "are", "was", "were", "with", "as", "by", "that", "this", "it",
}


def summarize(text: str, max_sentences: int = 5) -> str:
    """Extractive summary: score sentences by frequency of their
    non-trivial words, then return the top-scoring sentences in their
    original order."""

    text = text.strip()
    if not text:
        return ""

    sentences = [s.strip() for s in _SENTENCE_SPLIT_RE.split(text) if s.strip()]
    if len(sentences) <= max_sentences:
        return " ".join(sentences)

    word_freq: dict[str, int] = {}
    for sentence in sentences:
        for word in _WORD_RE.findall(sentence.lower()):
            if word in _COMMON_WORDS or len(word) < 3:
                continue
            word_freq[word] = word_freq.get(word, 0) + 1

    scored = []
    for i, sentence in enumerate(sentences):
        words = _WORD_RE.findall(sentence.lower())
        score = sum(word_freq.get(w, 0) for w in words) / max(len(words), 1)
        scored.append((i, score, sentence))

    top = sorted(scored, key=lambda x: x[1], reverse=True)[:max_sentences]
    top_in_order = sorted(top, key=lambda x: x[0])
    return " ".join(s for _, _, s in top_in_order)


def extract_key_points(text: str, max_points: int = 5) -> list[str]:
    """Return the highest-scoring individual sentences as bullet points."""

    summary = summarize(text, max_sentences=max_points)
    return [s.strip() for s in _SENTENCE_SPLIT_RE.split(summary) if s.strip()]
